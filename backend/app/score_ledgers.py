"""按记账快照查询逐人流水；MySQL 8 在数据库内筛选及分页。"""
import json
from datetime import timedelta
from pathlib import Path
from tempfile import SpooledTemporaryFile
from typing import Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select, text
from starlette.background import BackgroundTask

from . import auth
from .adjustments import month_at
from .db import get_db, now_utc
from .models import Department

router = APIRouter(prefix="/api/score-ledgers", tags=["score-ledgers"])
PAGE_SIZE = 20

# JSON_TABLE 将一批快照按记分员工展开，与唯一的 batch_id/employee_id 流水对应。
BASE = """
SELECT l.id, l.batch_id, l.employee_id, l.month, l.delta,
       l.before_total, l.final_total, l.override_total,
       p.person_name, p.avatar_url,
       JSON_EXTRACT(b.snapshot, CONCAT('$.rows[', p.row_index - 1, '].person.paths')) AS paths,
       JSON_UNQUOTE(JSON_EXTRACT(b.snapshot, '$.reason')) AS reason,
       JSON_UNQUOTE(JSON_EXTRACT(b.snapshot, '$.operator_name')) AS operator_name,
       b.created_at, c.cancelled_at
FROM orvibo_jifen_batches b
JOIN JSON_TABLE(b.snapshot, '$.rows[*]' COLUMNS (
    row_index FOR ORDINALITY,
    person_id VARCHAR(36) PATH '$.person.id',
    person_name VARCHAR(200) PATH '$.person.name',
    avatar_url TEXT PATH '$.person.avatar_url'
)) AS p ON TRUE
JOIN orvibo_jifen_ledger l ON l.batch_id = b.id
    AND l.employee_id = p.person_id COLLATE utf8mb4_unicode_ci
LEFT JOIN orvibo_jifen_ledger_cancellations c ON c.ledger_id = l.id
WHERE b.month = :month AND (:scope IN ('all', 'departments') OR (:scope = 'own' AND b.operator_id = :operator_id))
"""


def decode_paths(value):
    return json.loads(value) if isinstance(value, str) else value or []


def ledger_filter(user, month, q="", department_id="", operation="all", status="all"):
    scope = "all" if user.role == "owner" else user.ledger_scope
    params = {"month": month, "scope": scope, "operator_id": user.id,
              "q": q.strip(), "department_id": department_id}
    allowed = []
    if scope == "departments":
        for index, dept in enumerate(sorted(set(user.ledger_department_ids or []))):
            params[f"allowed_{index}"] = dept
            allowed.append(f"JSON_CONTAINS(JSON_EXTRACT(paths, '$[*].ids[*]'), JSON_QUOTE(:allowed_{index}))")
    scope_sql = (" OR ".join(allowed) or "FALSE") if scope == "departments" else "TRUE"
    conditions = ["LOCATE(LOWER(:q), LOWER(person_name)) > 0"]
    if department_id:
        conditions.append("JSON_CONTAINS(JSON_EXTRACT(paths, '$[*].ids[*]'), JSON_QUOTE(:department_id))")
    if operation != "all":
        conditions.append("delta > 0" if operation == "plus" else "delta < 0")
    if status != "all":
        conditions.append("cancelled_at IS NULL" if status == "valid" else "cancelled_at IS NOT NULL")
    return f"WITH original AS ({BASE}), scoped AS (SELECT * FROM original WHERE {scope_sql}), filtered AS (SELECT * FROM scoped WHERE {' AND '.join(conditions)})", params


def query_ledgers(db, user, month, q="", department_id="", operation="all", status="all", page=1):
    filtered, params = ledger_filter(user, month, q, department_id, operation, status)
    summary = db.execute(text(filtered + """
        SELECT COUNT(*) AS total,
               COALESCE(SUM(CASE WHEN cancelled_at IS NULL AND delta > 0 THEN delta ELSE 0 END), 0) AS plus_total,
               COALESCE(SUM(CASE WHEN cancelled_at IS NULL AND delta < 0 THEN delta ELSE 0 END), 0) AS minus_total
        FROM filtered
    """), params).mappings().one()
    total = int(summary["total"])
    page = min(page, max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE))
    records = db.execute(text(filtered + """
        SELECT * FROM filtered ORDER BY created_at DESC, id ASC LIMIT :limit OFFSET :offset
    """), {**params, "limit": PAGE_SIZE, "offset": (page - 1) * PAGE_SIZE}).mappings()
    items = []
    for record in records:
        row = dict(record)
        row["paths"] = decode_paths(row["paths"])
        row["cancelled"] = row["cancelled_at"] is not None
        for field in ("created_at", "cancelled_at"):
            row[field] = row[field].isoformat() + "Z" if row[field] else None
        items.append(row)

    # 只取去重后的历史路径，不把整个月的人员流水加载到应用内。
    historical_paths = db.execute(text(filtered + """
        SELECT CAST(paths AS CHAR) AS paths, MAX(created_at) AS latest
        FROM scoped GROUP BY CAST(paths AS CHAR) ORDER BY latest DESC, paths
    """), params).mappings()
    departments = {}
    current = {d.id: d for d in db.scalars(select(Department))}
    for record in historical_paths:
        for path in decode_paths(record["paths"]):
            ids, names = path["ids"], path["names"]
            if user.role != "owner" and user.ledger_scope == "departments" and not set(ids).intersection(user.ledger_department_ids or []):
                continue
            for index, department in enumerate(ids):
                departments.setdefault(department, {
                    "id": department, "name": names[index] + ("（已停用）" if department not in current or not current[department].active else ""), "parent_id": ids[index - 1] if index else "0",
                    "path_ids": ids[:index + 1], "path_names": names[:index + 1],
                })
    return {"items": items, "total": total, "page": page, "page_size": PAGE_SIZE, "month": month,
            "summary": {"plus_total": int(summary["plus_total"]), "minus_total": int(summary["minus_total"])},
            "departments": list(departments.values())}


def export_workbook(db, user, month, q="", department_id="", operation="all", status="all"):
    """共用页面的权限和筛选，逐批读取全部匹配流水，生成真正的 xlsx 文件。"""
    filtered, params = ledger_filter(user, month, q, department_id, operation, status)
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("积分明细")
    sheet.freeze_panes = "A2"
    headings = ["序号", "记分月份", "员工", "记分时部门", "加减分", "记分前总分", "记分后总分",
                "总分计算方式", "原因", "操作时间（北京时间）", "操作人", "有效状态", "作废时间（北京时间）"]
    for column, width in zip("ABCDEFGHIJKLM", [8, 12, 22, 48, 12, 14, 14, 14, 60, 24, 20, 12, 24]):
        sheet.column_dimensions[column].width = width

    def cell(value, *, header=False, cancelled=False):
        result = WriteOnlyCell(sheet, ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value)
        if isinstance(value, str):
            # 员工姓名、原因等始终按文本写入，避免被 Excel 当作公式执行。
            result.data_type = "s"
        result.alignment = Alignment(vertical="top", wrap_text=True)
        if header:
            result.font = Font(bold=True, color="FFFFFF")
            result.fill = PatternFill("solid", fgColor="53687B")
        elif cancelled:
            result.font = Font(color="808080")
        return result

    def beijing(value):
        return (value + timedelta(hours=8)).strftime("%Y-%m-%d %H:%M:%S") if value else ""

    try:
        sheet.append([cell(heading, header=True) for heading in headings])
        count = 0
        with db.execute(text(filtered + " SELECT * FROM filtered ORDER BY created_at DESC, id ASC"),
                        params, execution_options={"yield_per": 1000}) as records:
            for count, row in enumerate(records.mappings(), 1):
                if count >= 1_048_576:
                    raise HTTPException(413, "导出记录超过 Excel 行数限制，请缩小筛选范围")
                cancelled = row["cancelled_at"] is not None
                departments = "；".join(" / ".join(path["names"]) for path in decode_paths(row["paths"])) or "直属企业"
                values = [count, row["month"], row["person_name"], departments, row["delta"],
                          row["before_total"], row["final_total"],
                          "手动覆盖" if row["override_total"] is not None else "自动计算", row["reason"],
                          beijing(row["created_at"]), row["operator_name"], "已作废" if cancelled else "有效",
                          beijing(row["cancelled_at"])]
                cells = [cell(value, cancelled=cancelled) for value in values]
                cells[4].number_format = "+0;-0;0"
                sheet.append(cells)
        sheet.auto_filter.ref = f"A1:M{count + 1}"
        workbook.save(output)
        output.seek(0)
        return output
    except BaseException:
        output.close()
        raise
    finally:
        # write_only 工作表使用临时 XML 文件，生成失败时也需要清理。
        if not sheet.closed:
            sheet.close()
        if sheet._writer and Path(sheet._writer.out).exists():
            sheet._writer.cleanup()
        workbook.close()


def file_chunks(file):
    try:
        while chunk := file.read(64 * 1024):
            yield chunk
    finally:
        file.close()


@router.get("/export")
def export_score_ledgers(month: str | None = Query(None, pattern=r"^[0-9]{4}-(0[1-9]|1[0-2])$"),
                         q: str = Query("", max_length=200), department_id: str = Query("", max_length=100),
                         operation: Literal["all", "plus", "minus"] = "all",
                         status: Literal["all", "valid", "cancelled"] = "all",
                         user=Depends(auth.permission("ledgers.export")), db=Depends(get_db)):
    month = month or month_at(now_utc())
    file = export_workbook(db, user, month, q, department_id, operation, status)
    filename = quote(f"积分明细_{month}.xlsx")
    return StreamingResponse(file_chunks(file), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=score-ledgers-{month}.xlsx; filename*=UTF-8''{filename}"},
                             background=BackgroundTask(file.close))


@router.get("")
def score_ledgers(month: str | None = Query(None, pattern=r"^[0-9]{4}-(0[1-9]|1[0-2])$"),
                  q: str = Query("", max_length=200), department_id: str = Query("", max_length=100),
                  operation: Literal["all", "plus", "minus"] = "all",
                  status: Literal["all", "valid", "cancelled"] = "all",
                  page: int = Query(1, ge=1), user=Depends(auth.permission("ledgers.read")), db=Depends(get_db)):
    return query_ledgers(db, user, month or month_at(now_utc()), q, department_id, operation, status, page)
